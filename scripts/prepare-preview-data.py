"""Read one archived attachment and produce disposable replay preview data."""
import json
import os
import sys
import zipfile
from html import escape
from pathlib import Path
from xml.etree import ElementTree

# 只读预览的规模上限：整表会变成静态 HTML，太大就拖慢资料库页。
MAX_SHEET_ROWS = int(os.environ.get('CEOBE_SHEET_MAX_ROWS') or 500)
MAX_SHEET_COLS = int(os.environ.get('CEOBE_SHEET_MAX_COLS') or 60)

source, output = map(Path, sys.argv[1:3])
kind = sys.argv[3]
output.mkdir(parents=True, exist_ok=True)
data = {'pages': [], 'sheets': [], 'html': '', 'text': None, 'text_metadata': None}

if kind == 'pdf':
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(source)
    for index in range(len(pdf)):
        page = pdf[index]
        width, height = page.get_size()
        bitmap = page.render(scale=1.5)
        filename = f'pdf-page-{index + 1}.png'
        bitmap.to_pil().save(output / filename)
        data['pages'].append({'width': width, 'height': height, 'src': filename})
        bitmap.close()
        page.close()
    pdf.close()

def rgb(color, default):
    return '#' + color.rgb[-6:] if color is not None and color.type == 'rgb' else default

def display_value(cell):
    """按 Excel 的数字格式把值渲染成文本（只读预览用）。"""
    import datetime
    value = cell.value
    if value is None:
        return ''
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, (int, float)):
        fmt = (cell.number_format or 'General').lower()
        if '%' in fmt:
            decimals = fmt.count('0', fmt.find('.') + 1, fmt.find('%')) if '.' in fmt else 0
            return f'{value * 100:.{decimals}f}%'
        if ',' in fmt.split('.')[0]:
            decimals = fmt.count('0', fmt.find('.') + 1) if '.' in fmt else 0
            return f'{value:,.{decimals}f}'
        if '.' in fmt:
            decimals = sum(1 for ch in fmt.split('.')[1] if ch in '0#')
            if decimals:
                return f'{value:.{decimals}f}'
        return str(value)
    return str(value)

if kind == 'xlsx':
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter, column_index_from_string
    book = load_workbook(source, data_only=True)
    for sheet in book:
        total_rows = sheet.max_row or 0
        total_cols = sheet.max_column or 0
        rows = []
        for row in sheet:
            rows.append([{
                'value': display_value(cell),
                'bold': bool(cell.font.b), 'italic': bool(cell.font.i),
                'color': rgb(cell.font.color, '#111111'),
                'fill': rgb(cell.fill.fgColor, '#ffffff') if cell.fill.patternType == 'solid' else '#ffffff',
                'align': cell.alignment.horizontal or ('right' if isinstance(cell.value, (int, float)) else 'left'),
                'wrap': bool(cell.alignment.wrap_text),
                'size': float(cell.font.sz or 11) * 4 / 3,
            } for cell in row])
        # 只读预览要截断：整表写进静态片段，几千行就会拖慢资料库。
        rows = [row[:MAX_SHEET_COLS] for row in rows[:MAX_SHEET_ROWS]]
        truncated = {'totalRows': total_rows, 'totalCols': total_cols} if (total_rows > len(rows) or total_cols > MAX_SHEET_COLS) else None
        # 原表里被隐藏的行/列照常渲染但标灰，读者不会误以为它们没被隐藏。
        hidden_rows = [row for row, dimension in sheet.row_dimensions.items() if getattr(dimension, 'hidden', False)]
        hidden_cols = [column_index_from_string(letter) for letter, dimension in sheet.column_dimensions.items()
                       if getattr(dimension, 'hidden', False) and str(letter).isalpha()]
        columns = min(total_cols, MAX_SHEET_COLS)
        data['sheets'].append({
            'name': sheet.title, 'rows': rows,
            'widths': [round((sheet.column_dimensions[get_column_letter(c)].width or 13) * 7 + 5) for c in range(1, columns + 1)],
            'heights': [round((sheet.row_dimensions[r].height or sheet.sheet_format.defaultRowHeight or 15) * 4 / 3) for r in range(1, len(rows) + 1)],
            'merges': [[r.min_row, r.min_col, r.max_row, r.max_col] for r in sheet.merged_cells.ranges],
            'hiddenRows': hidden_rows, 'hiddenCols': hidden_cols,
            **({'truncated': truncated} if truncated else {}),
        })
    book.close()

if kind == 'txt':
    raw = source.read_bytes()
    encoding = 'utf-8-sig' if raw.startswith(b'\xef\xbb\xbf') else 'utf-8'
    try:
        text = raw.decode(encoding)
    except UnicodeDecodeError:
        encoding = 'utf-16' if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else 'gb18030'
        text = raw.decode(encoding, errors='replace')
    data['text'] = text
    data['text_metadata'] = {
        'bytes': len(raw), 'characters': len(text), 'lines': len(text.splitlines()),
        'encoding': encoding, 'replacement_characters': text.count('\ufffd'),
        'line_endings': {
            'crlf': text.count('\r\n'),
            'lf': text.count('\n') - text.count('\r\n'),
            'cr': text.count('\r') - text.count('\r\n'),
        },
    }
    data['html'] = '<pre class="ceobe-text-preview">' + escape(text) + '</pre>'

if kind == 'docx':
    with zipfile.ZipFile(source) as package:
        document = ElementTree.fromstring(package.read('word/document.xml'))
    namespace = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    word = '{' + namespace['w'] + '}'

    def on(element, name):
        node = element.find('w:' + name, namespace) if element is not None else None
        return node is not None and node.get(word + 'val', '1') not in {'0', 'false', 'off'}

    def css(values):
        return '; '.join(values) + (';' if values else '')

    def render_run(run):
        properties = run.find('w:rPr', namespace)
        styles = []
        if on(properties, 'b'):
            styles.append('font-weight: bold')
        if on(properties, 'i'):
            styles.append('font-style: italic')
        if on(properties, 'strike'):
            styles.append('text-decoration: line-through')
        underline = properties.find('w:u', namespace) if properties is not None else None
        if underline is not None and underline.get(word + 'val', 'single') not in {'none', '0', 'false'}:
            styles.append('text-decoration: underline')
        size = properties.find('w:sz', namespace) if properties is not None else None
        if size is not None and size.get(word + 'val', '').isdigit():
            styles.append(f"font-size: {int(size.get(word + 'val')) / 2:g}pt")
        color = properties.find('w:color', namespace) if properties is not None else None
        if color is not None and color.get(word + 'val') not in {None, 'auto'}:
            styles.append('#' + color.get(word + 'val')[-6:])
            styles[-1] = 'color: ' + styles[-1]
        language = properties.find('w:lang', namespace) if properties is not None else None
        lang = language.get(word + 'val') if language is not None else None
        content = []
        for child in run:
            if child.tag == word + 't':
                content.append(escape(child.text or ''))
            elif child.tag == word + 'tab':
                content.append('&#9;')
            elif child.tag in {word + 'br', word + 'cr'}:
                content.append('<br>')
        attrs = []
        if lang:
            attrs.append('lang="' + escape(lang, quote=True) + '"')
        run_css = css(styles)
        if run_css:
            attrs.append('style="' + run_css + '"')
        return '<span' + (' ' + ' '.join(attrs) if attrs else '') + '>' + ''.join(content) + '</span>'

    paragraphs = []
    for paragraph in document.findall('.//w:p', namespace):
        properties = paragraph.find('w:pPr', namespace)
        styles = []
        alignment = properties.find('w:jc', namespace) if properties is not None else None
        if alignment is not None:
            value = alignment.get(word + 'val')
            if value in {'left', 'right', 'center', 'justify'}:
                styles.append('text-align: ' + value)
        runs = []
        for node in paragraph.iter():
            if node.tag == word + 'r':
                runs.append(render_run(node))
        attrs = ' style="' + css(styles) + '"' if styles else ''
        paragraphs.append('<p' + attrs + '>' + ''.join(runs) + '</p>')
    data['html'] = ''.join(paragraphs)

(output / 'manifest.json').write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
print(f"Prepared {len(data['pages'])} PDF pages, {len(data['sheets'])} worksheets and {kind} document content.")
